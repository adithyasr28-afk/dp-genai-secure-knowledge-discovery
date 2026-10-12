"""
report/check_report.py

Automated consistency checks on the final report (run after report/finalize.py):
unresolved labels, internal markers, raw TeX, citation range, figure/table
order, blank pages, page-number fill, secrets and the remaining placeholders.

Run from the project root:  python report/check_report.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).resolve().parent
PDF, DOCX = HERE / "final_report.pdf", HERE / "final_report.docx"
problems, notes = [], []

pdf = pymupdf.open(str(PDF))
pages = [p.get_text() for p in pdf]
text = "\n".join(pages)
body_start = next(i for i, t in enumerate(pages) if "INDEX TERMS" in t and "ABSTRACT" in t)
body = "\n".join(pages[body_start:])

# 1. unresolved labels, internal markers, raw TeX
for pat, what in [(r"\{(fig|tab|eq):", "unresolved label"), (r"\[@", "unresolved citation"),
                  (r"\[(TEAM|VERIFY|CONFIRM)", "internal marker"), (r"\bTODO\b", "TODO"),
                  (r"\\(varepsilon|sqrt|frac|mathcal|approx|delta|rho)\b", "raw TeX"),
                  (r"Draft Status", "draft notes")]:
    hits = re.findall(pat, text)
    (problems if hits else notes).append(f"{what}: {len(hits)}")

# 1b. control characters in the sources (an escaped "\v", "\f" or "\a" in an editing
#     script silently turns "\varepsilon" into a vertical tab, etc.)
ctrl = []
for src in sorted((HERE / "src").glob("*.md")):
    raw = src.read_bytes()
    ctrl += [f"{src.name}@{m.start()}" for m in re.finditer(rb"[\x00-\x08\x0b\x0c\x0e-\x1f]|\r(?!\n)", raw)]
(problems if ctrl else notes).append(f"control characters in report/src: {ctrl}")

# 2. citations within the reference list
refs = re.findall(r"^\[(\d+)\] ", text, flags=re.M)
n_refs = max(map(int, refs)) if refs else 0
main_text = body.split("REFERENCES")[0]
cited = {int(x) for x in re.findall(r"\[(\d{1,2})\]", main_text)}
bad = sorted(c for c in cited if c < 1 or c > n_refs)
(problems if bad else notes).append(f"citations outside 1..{n_refs}: {bad}")
uncited = sorted(set(range(1, n_refs + 1)) - cited)
(problems if uncited else notes).append(f"references never cited in the body: {uncited}")

# 3. figure / table order and page fill
pm = json.loads((HERE / "page_map.json").read_text(encoding="utf-8"))
for kind in ("fig", "tab"):
    seq = [(k, v) for k, v in pm.items() if k.startswith(kind + ":")]
    pagesq = [v for _, v in seq]
    ok = pagesq == sorted(pagesq)
    (notes if ok else problems).append(f"{kind} captions in page order: {ok} ({len(seq)} items)")
(notes if not pm["_missing"] else problems).append(f"page map missing: {pm['_missing']}")

d = docx.Document(str(DOCX))
for ti, name in [(0, "Index"), (1, "List of Figures"), (2, "List of Tables")]:
    empty = [r.cells[1].text[:30] for r in d.tables[ti].rows[1:] if not r.cells[-1].text.strip()]
    (problems if empty else notes).append(f"{name}: rows without page number: {len(empty)} {empty[:3]}")

# 4. blank pages
blank = [i + 1 for i, p in enumerate(pdf) if len(pages[i].strip()) < 40 and not p.get_images()]
(problems if blank else notes).append(f"blank pages: {blank}")

# 5. secrets
sec = re.findall(r"AQ\.[A-Za-z0-9_-]{20,}|AIza[0-9A-Za-z_-]{30,}|sk-ant-api[0-9A-Za-z_-]+", text)
(problems if sec else notes).append(f"secret-like strings: {len(sec)}")

# 6. placeholders that are intentional (missing human input)
ph = set(re.findall(r"\[(?:Student name|Reg\. No\.[^\]]*|VIT e-mail|words)\]", text))
docx_text = "\n".join(p.text for p in d.paragraphs) + "\n".join(
    c.text for t in d.tables for r in t.rows for c in r.cells)
ph |= set(re.findall(r"\[(?:Student name|Reg\. No\.[^\]]*|VIT e-mail|words)\]", docx_text))
notes.append(f"intentional placeholders (need team input): {sorted(set(ph))}")
notes.append(f"pages: {len(pdf)}; body starts on page {body_start + 1}; references: {n_refs}")

print("PROBLEMS:" if problems else "PROBLEMS: none")
for p in problems:
    print("  -", p)
print("CHECKS:")
for n in notes:
    print("  -", n)
sys.exit(1 if problems else 0)
